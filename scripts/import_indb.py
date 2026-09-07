"""Import the INDB workbook into an isolated, validated SQLite reference database.
Run with Python + openpyxl. Does not modify the workbook or app database.
"""
import argparse
import hashlib
import json
import math
import os
from pathlib import Path
import sqlite3
import tempfile
import openpyxl

TEXT_COLUMNS = {"food_code", "food_name", "primarysource", "servings_unit"}
REQUIRED = {"food_code", "food_name", "primarysource", "energy_kcal", "servings_unit", "unit_serving_energy_kcal"}

def validate(headers, rows):
    if len(set(headers)) != len(headers) or not REQUIRED.issubset(headers):
        raise ValueError("Duplicate headers or missing required columns")
    if any(not isinstance(h, str) or not h.replace("_", "").isalnum() for h in headers):
        raise ValueError("Invalid column names")
    records, seen = [], set()
    for row_number, values in rows:
        if not any(v is not None for v in values):
            continue
        record = dict(zip(headers, values))
        for field in ("food_code", "food_name", "primarysource"):
            if not isinstance(record[field], str) or not record[field].strip():
                raise ValueError(f"Row {row_number}: missing {field}")
        if record["food_code"] in seen:
            raise ValueError(f"Duplicate food code: {record['food_code']}")
        seen.add(record["food_code"])
        for field, value in record.items():
            if field in TEXT_COLUMNS:
                if value is not None and not isinstance(value, str):
                    raise ValueError(f"Row {row_number}: invalid text {field}")
            elif value is not None and (isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0):
                raise ValueError(f"Row {row_number}: invalid nutrient {field}")
        if record["energy_kcal"] is None:
            raise ValueError(f"Row {row_number}: missing per-100g energy")
        issues = []
        if record["unit_serving_energy_kcal"] is None:
            issues.append("missing_serving_calories")
        if not (record["servings_unit"] or "").strip():
            issues.append("missing_serving_unit")
        record.update(source_row=row_number, missing_serving_calories=int("missing_serving_calories" in issues),
                      missing_serving_unit=int("missing_serving_unit" in issues), serving_usable=int(not issues))
        records.append(record)
    if not records:
        raise ValueError("No food records")
    return records

def run(source, output):
    source, output = Path(source), Path(output)
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=False)
    if len(workbook.sheetnames) != 1:
        raise ValueError("Expected one data sheet; inspect workbook before importing")
    sheet = workbook.active
    iterator = sheet.iter_rows()
    headers = [cell.value for cell in next(iterator)]
    raw = []
    for row in iterator:
        if any(cell.data_type == "f" for cell in row):
            raise ValueError("Formula found; resolve and verify before importing")
        raw.append((row[0].row, [cell.value for cell in row]))
    records = validate(headers, raw)
    metadata = {"dataset": "INDB 2024", "source_filename": source.name,
                "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                "source_sheet": sheet.title, "source_columns": headers,
                "nutrient_basis": "Unprefixed nutrient columns: per 100 g; unit_serving_ columns: per reference serving.",
                "methodology_url": "https://www.anuvaad.org.in/methodology/",
                "import_version": 1}
    workbook.close()
    output.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(suffix=".sqlite", dir=output)
    os.close(fd)
    try:
        db = sqlite3.connect(temporary)
        definitions = []
        for h in headers:
            kind = "TEXT" if h in TEXT_COLUMNS else "REAL"
            constraint = " PRIMARY KEY NOT NULL" if h == "food_code" else ""
            if h in {"food_name", "primarysource", "energy_kcal"}:
                constraint = " NOT NULL"
            if kind == "REAL":
                constraint += f' CHECK ("{h}" IS NULL OR "{h}" >= 0)'
            definitions.append(f'"{h}" {kind}{constraint}')
        definitions += ["source_row INTEGER NOT NULL", "missing_serving_calories INTEGER NOT NULL CHECK(missing_serving_calories IN (0,1))", "missing_serving_unit INTEGER NOT NULL CHECK(missing_serving_unit IN (0,1))", "serving_usable INTEGER NOT NULL CHECK(serving_usable IN (0,1))"]
        db.execute("CREATE TABLE nutrition_foods (" + ",".join(definitions) + ")")
        fields = list(records[0])
        columns = ",".join('"'+h+'"' for h in fields)
        db.executemany("INSERT INTO nutrition_foods ("+columns+") VALUES ("+",".join("?" for _ in fields)+")", [[r[h] for h in fields] for r in records])
        db.execute("CREATE TABLE nutrition_metadata (key TEXT PRIMARY KEY, value_json TEXT NOT NULL)")
        db.executemany("INSERT INTO nutrition_metadata VALUES (?,?)", [(k,json.dumps(v)) for k,v in metadata.items()])
        db.commit()
        assert db.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        # Compare every imported source field, including exact numeric values and nulls.
        loaded = db.execute("SELECT "+",".join('"'+h+'"' for h in headers)+" FROM nutrition_foods ORDER BY source_row").fetchall()
        assert loaded == [tuple(r[h] for h in headers) for r in records], "Source reconciliation failed"
        missing = [{"food_code":r["food_code"], "food_name":r["food_name"], "source_row":r["source_row"],
                    "issues":[k for k in ("missing_serving_calories", "missing_serving_unit") if r[k]]} for r in records if not r["serving_usable"]]
        report = {**metadata, "record_count":len(records), "source_column_count":len(headers),
                  "missing_serving_calories":sum(r["missing_serving_calories"] for r in records),
                  "missing_serving_unit":sum(r["missing_serving_unit"] for r in records),
                  "usable_reference_servings":sum(r["serving_usable"] for r in records),
                  "source_reconciliation":"Every source cell matches imported table", "sqlite_integrity":"ok",
                  "flagged_records":missing}
        db.close()
        os.replace(temporary, output/"indb-2024.sqlite")
        (output/"validation.json").write_text(json.dumps(report, indent=2, ensure_ascii=False)+"\n")
        print(json.dumps({k:report[k] for k in ("record_count","source_column_count","missing_serving_calories","missing_serving_unit","usable_reference_servings","sqlite_integrity")}))
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)

if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source")
    parser.add_argument("--output", default=str(Path(__file__).resolve().parents[1]/"data/nutrition"))
    args = parser.parse_args()
    run(args.source, args.output)
